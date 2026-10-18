"""
Conflicts importer for the Tabbycat API Importer v4/v5  (separate page: /conflicts)

ADDITION to the existing importer - nothing in app.py's existing code is changed. It reuses TabbycatAPI,
url_id, build_institution_index and find_existing_institution from app.py.

What it does
  * reads conflicts from a .csv / .xlsx (any sheet, any two columns, optional status column such as
    "Entered into Tab?" so rows marked Deleted are skipped)
  * "Mixed" mode: every row is classified on its own as judge-judge, judge-team (team name OR speaker name)
    or judge-institution; the two columns may be in either order
  * cleans typical sheet noise: "- Judge" suffixes, "Team " prefixes, "Name - Team name", "(nickname)",
    several entries in one cell (; , // newline), & vs and, accents, spacing, word order
  * batch-aware: conflicts already on the tab are skipped; same file can be uploaded again
  * rows that cannot be matched are listed with the reason and "did you mean" suggestions

Conflicts are saved with PATCH /tournaments/<slug>/adjudicators/<id> (fields institution_conflicts,
team_conflicts, adjudicator_conflicts).

Hook-up (2 lines at the bottom of app.py, above `if __name__ == '__main__':`):

    from conflicts import register_conflicts
    register_conflicts(app, TabbycatAPI, url_id, build_institution_index, find_existing_institution)
"""

import base64
import csv
import difflib
import io
import re
import time
import unicodedata

import requests
from flask import Blueprint, Response, render_template, request

TIME_BUDGET_SECONDS = 100          # Render free tier: gunicorn --timeout 120

TYPES = {
    'mixed': 'Mixed (auto-detect every row)',
    'adj_adj': 'Judge \u2194 judge only',
    'adj_inst': 'Judge \u2194 institution only',
    'team_adj': 'Judge \u2194 team / speaker only',
}
SLOT_OTHER = {'mixed': {'J', 'T', 'I'}, 'adj_adj': {'J'}, 'adj_inst': {'I'}, 'team_adj': {'T'}}
ROLE_KIND = {'J': 'adj_adj', 'T': 'team_adj', 'I': 'adj_inst'}
ROLE_NAME = {'J': 'judge', 'T': 'team or speaker', 'I': 'institution'}
KIND_FIELD = {'adj_adj': 'adjudicator_conflicts', 'adj_inst': 'institution_conflicts', 'team_adj': 'team_conflicts'}
FIELD_RESOURCE = {'adjudicator_conflicts': 'adjudicators', 'institution_conflicts': 'institutions',
                  'team_conflicts': 'teams'}

DEFAULT_BLOCKS = [
    {'n': 1, 'on': True, 'type': 'mixed', 'sheet': '', 'cols': 'A,B', 'status': 'C'},
    {'n': 2, 'on': True, 'type': 'adj_adj', 'sheet': '', 'cols': 'E,F', 'status': 'G'},
    {'n': 3, 'on': False, 'type': 'adj_inst', 'sheet': '', 'cols': 'A,B', 'status': ''},
]

SAMPLE_CSV = ("Judge or team or speaker,Judge or institution\n"
              "Jane Doe (ABC),John Roe - Judge\n"
              "Alex Speaker,Jane Doe\n"
              "Jane Doe,Example University\n")

HEADER_RE = re.compile(r'^(judge|judges|adjudicator|adjudicators|adj|team|teams|speaker|speakers|'
                       r'institution|institutions|university|name|conflict|conflicts|column|col)\b')
ROLE_TAIL = re.compile(r'\s*[-\u2013\u2014,]?\s*\(?\b(?:judge|adjudicator|adj)\b\)?\s*$', re.I)
LIST_SEP = re.compile(r'\s*(?:;|\r?\n|//|\\+|,)\s*')


# =============================================================================
# TEXT MATCHING
# =============================================================================

def loose(text):
    """
    Matching key: ignores case, accents, punctuation, spacing, '&' vs 'and' and a leading '[tournament] ' tag.
    'Vince  Hubilla (UPLB)' -> 'vince hubilla uplb'   '[abp2026] Ting & Rajaendran' -> 'ting and rajaendran'
    """
    s = str(text if text is not None else '').replace('\u0111', 'd').replace('\u0110', 'D')
    s = unicodedata.normalize('NFKD', s)
    s = ''.join(ch for ch in s if not unicodedata.combining(ch))
    s = re.sub(r'^\s*\[[^\]]*\]\s*', '', s)
    s = s.replace('&', ' and ')
    s = re.sub(r'[^\w\s]', ' ', s)
    return re.sub(r'\s+', ' ', s).strip().casefold()


def tok_key(key):
    return ' '.join(sorted(key.split()))


class Table:
    """name -> records, matched exactly (loose) first and, as a fallback, ignoring word order."""

    def __init__(self):
        self.exact, self.tokens, self.display = {}, {}, {}

    def add(self, text, rec, display=None):
        key = loose(text)
        if key:
            self.exact.setdefault(key, {})[rec['id']] = rec
            self.tokens.setdefault(tok_key(key), {})[rec['id']] = rec
            self.display.setdefault(key, display or str(text))

    def pick(self, text, token=False):
        key = loose(text)
        if not key:
            return ('none', None, None)
        found = list((self.tokens if token else self.exact).get(tok_key(key) if token else key, {}).values())
        if len(found) == 1:
            return ('ok', found[0], 'word order ignored' if token else None)
        return ('ambiguous', None, None) if found else ('none', None, None)

    def suggest(self, text, n=3):
        close = difflib.get_close_matches(loose(text), list(self.display), n=n, cutoff=0.78)
        return [self.display[k] for k in close]


# =============================================================================
# LOOKUP OF WHAT IS ON THE TAB SITE
# =============================================================================

class Lookup:
    def __init__(self, api, existing, inst_index, find_inst, url_id):
        self.api, self.inst_index, self.find_inst, self.url_id = api, inst_index, find_inst, url_id
        self.adjs, self.teams, self.speakers, self.insts = Table(), Table(), Table(), Table()
        self.team_recs = {}
        self.role_fn = {'J': self.judge, 'T': self.team_or_speaker, 'I': self.institution}

        for inst_id, info in inst_index['by_id'].items():
            rec = {'id': inst_id, 'url': self.build_url('institutions', inst_id),
                   'label': info.get('name') or info.get('code') or str(inst_id)}
            self.insts.add(info.get('name'), rec)
            self.insts.add(info.get('code'), rec)

        for a in existing['adjudicators']:
            a_id = a.get('id') if a.get('id') is not None else url_id(a.get('url'), 'adjudicators')
            if a_id is None or not a.get('name'):
                continue
            info = inst_index['by_id'].get(url_id(a.get('institution'), 'institutions'), {})
            code = info.get('code') or ''
            rec = {'id': a_id, 'raw': a, 'name': a['name'],
                   'url': (a.get('url') or '').rstrip('/') or self.build_url('adjudicators', a_id),
                   'label': f"{a['name']} ({code})" if code else a['name']}
            self.adjs.add(a['name'], rec, rec['label'])
            for extra in (info.get('code'), info.get('name')):
                if extra:
                    self.adjs.add(f"{a['name']} {extra}", rec, rec['label'])

        for t in existing['teams']:
            t_id = t.get('id') if t.get('id') is not None else url_id(t.get('url'), 'teams')
            if t_id is None:
                continue
            info = inst_index['by_id'].get(url_id(t.get('institution'), 'institutions'), {})
            ref, sref = t.get('reference'), t.get('short_reference')
            rec = {'id': t_id, 'url': (t.get('url') or '').rstrip('/') or self.build_url('teams', t_id),
                   'label': t.get('short_name') or t.get('long_name') or ref or str(t_id)}
            self.team_recs[t_id] = rec
            names = {t.get('short_name'), t.get('long_name'), ref, sref, t.get('code_name')}
            for prefix in (info.get('code'), info.get('name')):
                if prefix:
                    names.add(f"{prefix} {ref}")
                    if sref:
                        names.add(f"{prefix} {sref}")
            for n in names:
                self.teams.add(n, rec, rec['label'])

        for t in existing['teams']:                       # speakers nested in teams (if they carry names)
            t_id = t.get('id') if t.get('id') is not None else url_id(t.get('url'), 'teams')
            for sp in (t.get('speakers') or []):
                if isinstance(sp, dict) and sp.get('name') and t_id in self.team_recs:
                    self.speakers.add(sp['name'], self.team_recs[t_id])
        for sp in existing['speakers']:                   # /speakers endpoint
            if isinstance(sp, dict) and sp.get('name'):
                t_id = url_id(sp.get('team'), 'teams')
                if t_id in self.team_recs:
                    self.speakers.add(sp['name'], self.team_recs[t_id])

    def build_url(self, resource, item_id):
        if resource == 'institutions':
            return f"{self.api.base_url}/api/v1/institutions/{item_id}"
        return f"{self.api.base_url}/api/v1/tournaments/{self.api.slug}/{resource}/{item_id}"

    # ---- name -> record, per role ------------------------------------------------------------
    def judge(self, text):
        for token in (False, True):
            res = self.adjs.pick(text, token)
            if res[0] != 'none':
                return res
        return ('none', None, None)

    def team_or_speaker(self, text):
        for token in (False, True):                       # exact on team/speaker first, then word order ignored
            for table, tag in ((self.teams, None), (self.speakers, 'speaker')):
                st, rec, via = table.pick(text, token)
                if st == 'ok':
                    return ('ok', rec, ', '.join(x for x in (tag, via) if x) or None)
                if st == 'ambiguous':
                    return ('ambiguous', None, None)
        return ('none', None, None)

    def institution(self, text):
        found = self.find_inst(self.inst_index, text, text)
        if found:
            info = self.inst_index['by_id'].get(found[0], {})
            return ('ok', {'id': found[0] if found[0] is not None else found[1], 'url': found[1],
                           'label': info.get('name') or str(text)}, None)
        for token in (False, True):
            res = self.insts.pick(text, token)
            if res[0] != 'none':
                return res
        return ('none', None, None)

    def suggest(self, text, roles):
        tables = {'J': [self.adjs], 'T': [self.teams, self.speakers], 'I': [self.insts]}
        out = []
        for role in roles:
            for table in tables[role]:
                for s in table.suggest(text):
                    if s not in out:
                        out.append(s)
        return out[:3]

    # ---- conflicts already stored on an adjudicator -------------------------------------------
    def current_ids(self, rec, field):
        """Set of ids already conflicted, or None if the API does not report this field."""
        if not isinstance(rec['raw'].get(field), list):
            ok, data = api_get(self.api, rec['url'])         # list view lacked it: read the single record
            if ok and isinstance(data, dict):
                rec['raw'] = data
        lst = rec['raw'].get(field)
        if not isinstance(lst, list):
            return None
        return {self.url_id(x, FIELD_RESOURCE[field]) for x in lst} - {None}

    def as_url(self, item, resource):
        if isinstance(item, str) and item.startswith('http'):
            return item.rstrip('/')
        item_id = self.url_id(item, resource)
        return self.build_url(resource, item_id) if item_id is not None else item


# =============================================================================
# CELL CLEANING / RESOLVING
# =============================================================================

def _judge_tail(t):
    return ROLE_TAIL.sub('', t.strip()).strip()


def _no_team(t):
    return re.sub(r'^\s*team\s+', '', t, flags=re.I).strip()


def _unique(seq):
    out = []
    for x in seq:
        if x and x not in out:
            out.append(x)
    return out


def alt_variants(t):
    """Alternative readings of one entry: 'Name - Team X', 'Name from Team X', 'Name (nick)', 'A/B' ..."""
    t = _judge_tail(t)
    raw = []
    no_paren = re.sub(r'\([^)]*\)', ' ', t).strip()
    if no_paren != t:
        raw.append(no_paren)
    raw += re.findall(r'\(([^)]+)\)', t)
    for sep in (r'\s+[-\u2013\u2014]\s+', r'\s+from\s+', r'\s*/\s*', r'\s*:\s*'):
        parts = re.split(sep, t, flags=re.I)
        if len(parts) > 1:
            raw += parts
    out = []
    for v in raw:
        v = _judge_tail(v)
        out += [v, _no_team(v)]
    return _unique(out)


def gather(variants, allowed, lookup):
    found, amb = {}, set()
    for v in variants:
        if not loose(v):
            continue
        for role in allowed:
            st, rec, via = lookup.role_fn[role](v)
            if st == 'ok':
                found.setdefault(role, {})[rec['id']] = (rec, via, v)
            elif st == 'ambiguous':
                amb.add(role)
    return found, amb


def resolve_piece(text, allowed, lookup, alt):
    """-> (item, None) or (None, problem). item = {'text', 'roles': {role: (record, via, matched_text)}}."""
    t = text.strip()
    j = _judge_tail(t)
    found, amb = gather(_unique([j, _no_team(j)]), allowed, lookup)
    if not found and alt:
        found, amb2 = gather(alt_variants(t), allowed, lookup)
        amb |= amb2
    if not found:
        if not alt:
            return None, None
        if amb:
            return None, (f"'{text}' matches more than one {'/'.join(ROLE_NAME[r] for r in sorted(amb))} - "
                          f"make it more specific, e.g. 'Name (CODE)'")
        hint = lookup.suggest(j or t, allowed)
        return None, f"'{text}' not found" + (f" - did you mean: {' | '.join(hint)}?" if hint else '')
    roles = {}
    for role, d in found.items():
        if len(d) > 1:
            return None, f"'{text}' points to different {ROLE_NAME[role]}s: " + ' / '.join(v[0]['label'] for v in d.values())
        roles[role] = next(iter(d.values()))
    return {'text': text, 'roles': roles}, None


def resolve_cell(text, allowed, lookup):
    """A cell may hold one entry or a list (; , // newline). -> (items, problems)."""
    text = str(text).strip()
    item, _ = resolve_piece(text, allowed, lookup, alt=False)       # whole cell exactly as typed
    if item:
        return [item], []
    parts = [p for p in LIST_SEP.split(text) if p.strip()]
    if len(parts) > 1:
        items, problems, seen = [], [], set()
        for p in parts:
            it, why = resolve_piece(p, allowed, lookup, alt=True)
            if not it:
                problems.append(why)
                continue
            key = frozenset((r, v[0]['id']) for r, v in it['roles'].items())
            if key not in seen:
                seen.add(key)
                items.append(it)
        return items, problems
    item, why = resolve_piece(text, allowed, lookup, alt=True)
    return ([item], []) if item else ([], [why])


# =============================================================================
# FILE READING
# =============================================================================

def col_index(token):
    token = token.strip()
    if token.isdigit():
        return int(token) - 1
    n = 0
    for ch in token.upper():
        n = n * 26 + (ord(ch) - 64)
    return n - 1


def parse_cols(spec):
    toks = re.findall(r'[A-Za-z]+|\d+', spec or '')
    if len(toks) != 2:
        raise ValueError(f"columns must be two letters such as 'A,B' (got '{spec}')")
    return col_index(toks[0]), col_index(toks[1])


def load_table(raw, filename, sheet=''):
    """Whole sheet / csv as a list of rows (lists of stripped strings)."""
    name = (filename or '').lower()
    if name.endswith(('.xlsx', '.xlsm')):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        if sheet:
            match = next((s for s in wb.sheetnames if s.strip().casefold() == sheet.strip().casefold()), None)
            if match is None:
                raise ValueError(f"sheet '{sheet}' not found (sheets: {', '.join(wb.sheetnames)})")
            ws = wb[match]
        else:
            ws = wb.worksheets[0]
        return [[('' if c is None else str(c).strip()) for c in r] for r in ws.iter_rows(values_only=True)]
    text = None
    for enc in ('utf-8-sig', 'cp1252'):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        text = raw.decode('latin-1')
    first = next((ln for ln in text.splitlines() if ln.strip()), '')
    delim = ',' if ',' in first else ('\t' if '\t' in first else (';' if ';' in first else ','))
    return [[c.strip() for c in r] for r in csv.reader(io.StringIO(text), delimiter=delim)]


def read_block(table, cols, status_col, skip_statuses):
    """-> (pairs [(row, a, b)], problems [(row, a, b, why)], rows skipped because of their status)."""
    i1, i2 = cols
    si = col_index(status_col) if status_col else None
    pairs, problems, skipped, leading = [], [], 0, True
    for rownum, row in enumerate(table, start=1):
        def cell(i):
            return row[i].strip() if i is not None and 0 <= i < len(row) else ''
        a, b = cell(i1), cell(i2)
        if not (a or b):
            continue
        if leading:                                       # skip title / header rows ("Team Name", "Judge Name")
            if all(HEADER_RE.match(loose(c)) for c in (a, b) if c):
                continue
            leading = False
        if si is not None and loose(cell(si)) in skip_statuses:
            skipped += 1
            continue
        if not a or not b:
            problems.append((rownum, a, b, 'needs a value in both columns'))
            continue
        pairs.append((rownum, a, b))
    return pairs, problems, skipped


# =============================================================================
# HTTP HELPERS (TabbycatAPI only knows GET / POST, so PATCH / PUT live here)
# =============================================================================

def _headers(api):
    headers = {}
    if api.auth_method == 'session':
        csrf = api.session.cookies.get('csrftoken')
        if csrf:
            headers['X-CSRFToken'] = csrf
            headers['Referer'] = api.base_url + '/'
    return headers


def api_get(api, url):
    try:
        time.sleep(0.2)
        resp = api.session.get(url, timeout=30)
        if resp.status_code == 200:
            return True, resp.json()
        return False, f"HTTP {resp.status_code}"
    except (requests.exceptions.RequestException, ValueError) as e:
        return False, str(e)[:150]


def api_send(api, method, url, payload, retries=3):
    for attempt in range(retries):
        try:
            time.sleep(0.4)
            resp = api.session.request(method, url, json=payload, headers=_headers(api), timeout=30)
            if resp.status_code in (200, 201):
                try:
                    return True, resp.json()
                except ValueError:
                    return True, {}
            if resp.status_code == 429:
                time.sleep(2 ** attempt)
                continue
            return False, f"HTTP {resp.status_code} on {method} {url.replace(api.base_url, '')}: {resp.text[:300]}"
        except requests.exceptions.RequestException as e:
            if attempt == retries - 1:
                return False, f"Request failed on {method} {url.replace(api.base_url, '')}: {str(e)[:200]}"
            time.sleep(1)
    return False, 'rate limited (HTTP 429)'


# =============================================================================
# PROCESSING
# =============================================================================

def _roles_text(items):
    names = sorted({ROLE_NAME[r] for it in items for r in it['roles']})
    return '/'.join(names) or 'nothing'


def _note(role, via, matched, target_label):
    if role == 'T' and via and via.startswith('speaker'):
        return f"  [speaker '{matched}' \u2192 team {target_label}" + (", word order ignored]" if 'word' in via else "]")
    if via:
        return f"  [{via}: '{matched}']"
    return ''


def process_block(btype, pairs, lookup, pending, seen, section):
    other_allowed = SLOT_OTHER[btype]
    allowed = other_allowed | {'J'}
    for rownum, c1, c2 in pairs:
        row_label = f"Row {rownum}: {c1} \u2194 {c2}"
        items1, probs1 = resolve_cell(c1, allowed, lookup)
        items2, probs2 = resolve_cell(c2, allowed, lookup)
        if not items1 or not items2:
            section['failed'].append((row_label, '; '.join(probs1 + probs2) or 'nothing recognised', c1, c2))
            continue
        for why in probs1 + probs2:                       # list cells: the parts that matched still import
            section['failed'].append((row_label, f"part not imported - {why}", c1, c2))

        def judge_cell(items):
            return all('J' in it['roles'] for it in items)

        def other_cell(items):
            return all(any(r in it['roles'] for r in other_allowed) for it in items)

        v1 = judge_cell(items1) and other_cell(items2)
        v2 = judge_cell(items2) and other_cell(items1)
        if v1 and v2:
            if all(set(it['roles']) == {'J'} for it in items1 + items2):
                v2 = False                                # judge-judge: order does not matter
            else:
                section['failed'].append((row_label, f"ambiguous - both columns could be the judge "
                                          f"(column 1: {_roles_text(items1)}, column 2: {_roles_text(items2)})",
                                          c1, c2))
                continue
        if not v1 and not v2:
            r1 = {r for it in items1 for r in it['roles']}
            r2 = {r for it in items2 for r in it['roles']}
            if 'J' not in r1 | r2 and (('T' in r1 and 'I' in r2) or ('I' in r1 and 'T' in r2)):
                section['failed'].append((row_label, "team/speaker \u2194 institution is a team-institution conflict "
                                          "- not supported by this importer yet", c1, c2))
                continue
            section['failed'].append((row_label, f"no judge + {'/'.join(ROLE_NAME[r] for r in sorted(other_allowed))} "
                                      f"combination: column 1 = {_roles_text(items1)}, "
                                      f"column 2 = {_roles_text(items2)}", c1, c2))
            continue
        judges, others = (items1, items2) if v1 else (items2, items1)

        for o in others:
            opts = [r for r in sorted(other_allowed) if r in o['roles']]
            if len(opts) != 1:
                section['failed'].append((row_label, f"'{o['text']}' could be a "
                                          f"{' or a '.join(ROLE_NAME[r] for r in opts)} - make it more specific",
                                          c1, c2))
                continue
            role = opts[0]
            kind, field = ROLE_KIND[role], KIND_FIELD[ROLE_KIND[role]]
            target, via, matched = o['roles'][role]
            for jitem in judges:
                adj = jitem['roles']['J'][0]
                label = (f"Row {rownum}: {adj['label']} \u2194 {target['label']} ({ROLE_NAME[role].split(' ')[0]})"
                         + _note(role, via, matched, target['label']))
                if kind == 'adj_adj' and adj['id'] == target['id']:
                    section['failed'].append((label, 'both sides are the same judge', c1, c2))
                    continue
                dedupe = (kind, frozenset((adj['id'], target['id']))) if kind == 'adj_adj' \
                    else (kind, adj['id'], target['id'])
                if dedupe in seen:
                    section['existing'].append(label + '  (duplicate row in file)')
                    continue
                seen.add(dedupe)
                have = lookup.current_ids(adj, field)
                if have is None:
                    section['failed'].append((label, f"this API does not report '{field}' for the judge, skipped "
                                              f"to avoid overwriting existing conflicts", c1, c2))
                    continue
                already = target['id'] in have
                if kind == 'adj_adj' and not already:     # pair may be stored the other way round
                    back = lookup.current_ids(target, field)
                    already = bool(back) and adj['id'] in back
                if already:
                    section['existing'].append(label)
                    continue
                group = pending.setdefault(adj['id'], {'rec': adj, 'fields': {}, 'items': []})
                group['fields'].setdefault(field, {})[target['id']] = target['url']
                group['items'].append((section, label))


def apply_pending(api, lookup, pending, dry_run, deadline, notes):
    for grp in pending.values():
        rec = grp['rec']
        payload = {}
        for field, targets in grp['fields'].items():
            resource = FIELD_RESOURCE[field]
            current = [lookup.as_url(x, resource) for x in (rec['raw'].get(field) or [])]
            have = {lookup.url_id(u, resource) for u in current}
            payload[field] = current + [u for tid, u in targets.items() if tid not in have]

        def finish(ok, reason=''):
            for section, label in grp['items']:
                if ok:
                    section['created'].append(label)
                else:
                    section['failed'].append((label, reason, '', ''))

        if dry_run:
            finish(True)
            continue
        if time.time() > deadline:
            finish(False, 'ran out of time on this request - upload the same file again (conflicts already '
                          'added are skipped)')
            continue
        ok, data = api_send(api, 'PATCH', rec['url'], payload)
        if not ok and 'HTTP 405' in str(data):           # PATCH not allowed: fall back to a full PUT
            got, detail = api_get(api, rec['url'])
            if got and isinstance(detail, dict):
                detail.update(payload)
                ok, data = api_send(api, 'PUT', rec['url'], detail)
        if not ok:
            finish(False, str(data))
            continue
        finish(True)
        if isinstance(data, dict):                       # sanity check: did the site keep what we sent?
            for field, urls in payload.items():
                saved = data.get(field)
                if isinstance(saved, list):
                    resource = FIELD_RESOURCE[field]
                    want = {lookup.url_id(u, resource) for u in urls}
                    got_ids = {lookup.url_id(x, resource) for x in saved}
                    if not want <= got_ids:
                        notes.append(f"{rec['label']}: the site accepted the request but its reply did not list "
                                     f"every {field.replace('_', ' ')} sent - please check on the tab site.")
                        break


def failed_csv(failed):
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(['entry', 'column 1', 'column 2', 'problem'])
    for label, why, c1, c2 in failed:
        w.writerow([label, c1, c2, why])
    return base64.b64encode(('\ufeff' + out.getvalue()).encode('utf-8')).decode('ascii')


# =============================================================================
# FLASK ROUTES
# =============================================================================

def register_conflicts(app, TabbycatAPI, url_id, build_institution_index, find_existing_institution):
    bp = Blueprint('conflicts_importer', __name__)

    def base_ctx():
        return {'results': None, 'error': None, 'notes': [], 'dry_run': False, 'stats': None,
                'types': TYPES, 'blocks': [dict(b) for b in DEFAULT_BLOCKS], 'skip_status': 'deleted',
                'form': {'api_url': '', 'tournament_slug': '', 'api_username': ''}}

    @bp.route('/conflicts', methods=['GET', 'POST'])
    def conflicts_page():
        ctx = base_ctx()
        if request.method == 'GET':
            return render_template('conflicts.html', **ctx)

        started = time.time()
        deadline = started + TIME_BUDGET_SECONDS
        f = request.form
        ctx['form'] = {k: f.get(k, '').strip() for k in ('api_url', 'tournament_slug', 'api_username')}
        token, password = f.get('api_token', '').strip(), f.get('api_password', '')
        ctx['dry_run'] = dry_run = bool(f.get('dry_run'))
        ctx['skip_status'] = f.get('skip_status', 'deleted')
        skip_set = {loose(s) for s in re.split(r'[,;]', ctx['skip_status']) if s.strip()}
        ctx['blocks'] = [{'n': b['n'], 'on': bool(f.get(f"b{b['n']}_on")),
                          'type': f.get(f"b{b['n']}_type", b['type']),
                          'sheet': f.get(f"b{b['n']}_sheet", '').strip(),
                          'cols': f.get(f"b{b['n']}_cols", b['cols']).strip(),
                          'status': f.get(f"b{b['n']}_status", '').strip()} for b in DEFAULT_BLOCKS]

        def fail(msg):
            ctx['error'] = msg
            return render_template('conflicts.html', **ctx)

        if not ctx['form']['api_url'] or not ctx['form']['tournament_slug']:
            return fail('Enter the Tabbycat URL and the tournament slug.')
        if not token and not (ctx['form']['api_username'] and password):
            return fail('Enter an API token (or an admin username + password).')

        # ---- read the files -------------------------------------------------------------------
        shared = request.files.get('shared_file')
        shared_raw = shared.read() if shared and shared.filename else None
        tables, parsed, results = {}, [], []
        for b in ctx['blocks']:
            if not b['on']:
                continue
            title = f"Block {b['n']}: {TYPES.get(b['type'], b['type'])} (columns {b['cols']}" + \
                    (f", status {b['status']}" if b['status'] else '') + ')'
            own = request.files.get(f"b{b['n']}_file")
            if own and own.filename:
                raw, fname, src = own.read(), own.filename, f"b{b['n']}"
            elif shared_raw is not None:
                raw, fname, src = shared_raw, shared.filename, 'shared'
            else:
                return fail(f"Block {b['n']} is ticked but no file was uploaded for it.")
            if b['type'] not in TYPES:
                return fail(f"Block {b['n']}: unknown type.")
            try:
                cols = parse_cols(b['cols'])
                key = (src, b['sheet'])
                if key not in tables:
                    tables[key] = load_table(raw, fname, b['sheet'])
                pairs, problems, skipped = read_block(tables[key], cols, b['status'], skip_set)
            except Exception as e:
                return fail(f"{title}: {e}")
            section = {'title': title, 'rows': len(pairs) + len(problems), 'skipped': skipped, 'created': [],
                       'existing': [], 'failed': [(f"Row {n}", why, a, c) for n, a, c, why in problems]}
            parsed.append((b['type'], pairs, section))
            results.append(section)
        if not parsed:
            return fail('Tick at least one block and upload a file.')

        # ---- read the tab site --------------------------------------------------------------
        try:
            api = TabbycatAPI(ctx['form']['api_url'], token, ctx['form']['tournament_slug'],
                              ctx['form']['api_username'] or None, password or None)
            need_teams = any(t in ('mixed', 'team_adj') for t, _, _ in parsed)
            existing, problems = api.load_existing(need_teams=need_teams, need_adjudicators=True,
                                                   need_speakers=need_teams)
        except Exception as e:
            return fail(f"Could not connect to the tab site: {e}")
        if problems:
            return fail('Could not read the existing entries from the tab site (' + '; '.join(problems) +
                        '). Nothing was imported. Check the API token has admin access and the slug is right.')

        inst_index = build_institution_index(existing['institutions'], api.base_url)
        lookup = Lookup(api, existing, inst_index, find_existing_institution, url_id)

        pending, seen, notes = {}, set(), []
        for btype, pairs, section in parsed:
            process_block(btype, pairs, lookup, pending, seen, section)
        apply_pending(api, lookup, pending, dry_run, deadline, notes)

        for s in results:
            s['failed_csv'] = failed_csv(s['failed']) if s['failed'] else ''
        ctx.update(results=results, notes=notes, stats={'seconds': round(time.time() - started, 1)})
        return render_template('conflicts.html', **ctx)

    @bp.route('/conflicts/sample.csv')
    def conflicts_sample():
        return Response(SAMPLE_CSV, mimetype='text/csv',
                        headers={'Content-Disposition': 'attachment; filename=conflicts_sample.csv'})

    app.register_blueprint(bp)
