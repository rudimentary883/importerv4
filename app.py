"""
Tabbycat API Importer v5.0 - adds the combined "Teams + Speakers" file (one row = team + speakers laid out horizontally,
  standard layout OR Google-Form style headers) on top of the v4.0 batch-aware importer.
Tabbycat API Importer v4.0 — batch-aware ("dynamic") importing
  * reads the institutions / teams / judges / speakers ALREADY on the tab site
  * only imports entries that are NEW (safe to re-upload the same, updated .csv files)
  * institutions file is optional: teams / judges can use institutions already on the site
  * optional "preview only" mode that imports nothing
v3.6 base: speaker team URLs, category support, independent teams, institution regions
+ WSDC support (3-5 speakers per team)
"""

import os
import io
import csv
import json
import time
import re
import uuid
import requests
from collections import OrderedDict

from flask import Flask, render_template, request, send_file, flash, redirect, url_for, jsonify

from combined_import import parse_combined, inst_key, template_csv

app = Flask(__name__)
app.secret_key = os.environ.get('SECRET_KEY', 'tabbycat-importer-key-2024')
app.config['MAX_CONTENT_LENGTH'] = 16 * 1024 * 1024

# Max / min speakers imported per team, by debate format
FORMAT_MAX_SPEAKERS = {'bp': 2, '3v3': 3, 'wsdc': 5}
FORMAT_MIN_SPEAKERS = {'wsdc': 3}
FORMAT_LABELS = {'bp': 'BP', '3v3': '3v3', 'wsdc': 'WSDC'}


# Finished CSV downloads (kept server-side; cookies are far too small for this)
DOWNLOAD_CACHE = OrderedDict()
MAX_CACHED_DOWNLOADS = 8


def norm(value):
    """Normalise text for matching: trim, collapse spaces, ignore case."""
    return re.sub(r'\s+', ' ', str(value if value is not None else '')).strip().casefold()


def url_id(value, kind):
    """Numeric id from an API hyperlink ('.../teams/42' -> 42). Accepts ints, digit strings, dicts."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, dict):
        if value.get('id') is not None:
            return url_id(value.get('id'), kind)
        return url_id(value.get('url'), kind)
    if isinstance(value, int):
        return value
    text = str(value).strip().rstrip('/')
    if text.isdigit():
        return int(text)
    match = re.search(r'/%s/(\d+)' % re.escape(kind), text)
    return int(match.group(1)) if match else None


def clean_string(val):
    if val is None:
        return ''
    s = str(val).strip()
    return s if s != 'None' else ''


def parse_bool(val):
    if val is None:
        return False
    if isinstance(val, bool):
        return val
    return str(val).strip().upper() in ('TRUE', '1', 'YES', 'Y', 'T')


def parse_float_or_none(val):
    if val is None or str(val).strip() == '':
        return None
    try:
        return float(val)
    except (ValueError, TypeError):
        return None


def read_csv_file(file):
    text = file.read().decode('utf-8')
    reader = csv.DictReader(io.StringIO(text))
    return list(reader)


def read_excel_file(file):
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ImportError("openpyxl is required for Excel files")
    wb = load_workbook(file)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        row_dict = {}
        for i, header in enumerate(headers):
            if header:
                row_dict[header] = row[i] if i < len(row) else None
        rows.append(row_dict)
    return rows


def read_uploaded_file(file):
    ext = file.filename.rsplit('.', 1)[1].lower()
    if ext == 'csv':
        return read_csv_file(file)
    return read_excel_file(file)


# =============================================================================
# TABBYCAT API CLIENT (v3.6)
# =============================================================================

class TabbycatAPI:
    def __init__(self, base_url, token, tournament_slug, username=None, password=None):
        self.base_url = base_url.rstrip('/')
        self.token = token.strip() if token else ''
        self.slug = tournament_slug.strip('/')
        self.username = username
        self.password = password
        self.session = requests.Session()
        self.session.headers.update({
            'User-Agent': 'TabbycatImporter/3.6 (Render; Python requests)',
            'Accept': 'application/json',
            'Content-Type': 'application/json'
        })
        self.created_institutions = {}
        self.stats = {'success': 0, 'failed': 0, 'errors': []}
        self.auth_method = None
        self._authenticate()

    def _authenticate(self):
        if self.token:
            self.session.headers['Authorization'] = f'Token {self.token}'
            if self._test_auth():
                self.auth_method = 'token'
                return

        if self.username and self.password:
            self.session.headers.pop('Authorization', None)
            if self._login_session():
                self.auth_method = 'session'
                return

        if self.token:
            self.session.headers['Authorization'] = f'Token {self.token}'

    def _login_session(self):
        try:
            login_url = f"{self.base_url}/accounts/login/"
            resp = self.session.get(login_url, timeout=15)
            csrf_match = re.search(r'name="csrfmiddlewaretoken" value="([^"]+)"', resp.text)
            csrf_token = csrf_match.group(1) if csrf_match else ''

            login_data = {
                'username': self.username,
                'password': self.password,
                'csrfmiddlewaretoken': csrf_token,
                'next': '/'
            }
            resp = self.session.post(login_url, data=login_data, timeout=15)

            test_url = f"{self.base_url}/database/"
            resp = self.session.get(test_url, timeout=15)
            return resp.status_code == 200
        except Exception:
            return False

    def _test_auth(self):
        try:
            url = f"{self.base_url}/api/v1/institutions"
            resp = self.session.get(url, timeout=10)
            return resp.status_code in (200, 401)
        except Exception:
            return False

    def _global_url(self, path):
        return f"{self.base_url}/api/v1{path}"

    def _tournament_url(self, path):
        return f"{self.base_url}/api/v1/tournaments/{self.slug}{path}"

    def _request(self, method, url, data=None, retries=3):
        for attempt in range(retries):
            try:
                time.sleep(0.4)
                if method == 'POST':
                    resp = self.session.post(url, json=data, timeout=30)
                else:
                    resp = self.session.get(url, timeout=30)

                if resp.status_code in (200, 201):
                    self.stats['success'] += 1
                    return resp.json()
                elif resp.status_code == 429:
                    time.sleep(2 ** attempt)
                    continue
                else:
                    error = f"HTTP {resp.status_code} on {method} {url.replace(self.base_url, '')}: {resp.text[:300]}"
                    self.stats['errors'].append(error)
                    self.stats['failed'] += 1
                    return None
            except requests.exceptions.RequestException as e:
                if attempt == retries - 1:
                    error = f"Request failed on {method} {url.replace(self.base_url, '')}: {str(e)}"
                    self.stats['errors'].append(error)
                    self.stats['failed'] += 1
                    return None
                time.sleep(1)
        return None

    def test_connection(self):
        diagnostics = {
            'ok': False,
            'auth_method': self.auth_method,
            'steps': [],
            'suggestion': ''
        }

        try:
            resp = self.session.get(self.base_url, timeout=10, allow_redirects=True)
            diagnostics['steps'].append({
                'step': 'Base URL reachable',
                'status': resp.status_code,
                'ok': resp.status_code < 500
            })
        except Exception as e:
            diagnostics['steps'].append({
                'step': 'Base URL reachable',
                'status': 0,
                'ok': False,
                'error': str(e)
            })
            diagnostics['suggestion'] = 'Cannot reach your Tabbycat URL. Check for typos.'
            return diagnostics

        try:
            url = f"{self.base_url}/api/v1/institutions"
            resp = self.session.get(url, timeout=10)
            diagnostics['steps'].append({
                'step': 'Global institutions list (GET /api/v1/institutions)',
                'status': resp.status_code,
                'ok': resp.status_code == 200,
                'body_preview': resp.text[:100] if resp.text else ''
            })
        except Exception as e:
            diagnostics['steps'].append({
                'step': 'Global institutions list',
                'status': 0,
                'ok': False,
                'error': str(e)
            })

        try:
            url = f"{self.base_url}/api/v1/tournaments/{self.slug}/institutions"
            resp = self.session.get(url, timeout=10)
            diagnostics['steps'].append({
                'step': 'Tournament institutions list (GET)',
                'status': resp.status_code,
                'ok': resp.status_code == 200,
                'body_preview': resp.text[:100] if resp.text else ''
            })
        except Exception as e:
            diagnostics['steps'].append({
                'step': 'Tournament institutions list',
                'status': 0,
                'ok': False,
                'error': str(e)
            })

        try:
            url = f"{self.base_url}/api/v1/tournaments/{self.slug}/teams"
            resp = self.session.get(url, timeout=10)
            diagnostics['steps'].append({
                'step': 'Tournament teams list (GET)',
                'status': resp.status_code,
                'ok': resp.status_code in (200, 401),
                'body_preview': resp.text[:100] if resp.text else ''
            })

            if resp.status_code == 200:
                diagnostics['ok'] = True
                diagnostics['suggestion'] = 'Connection successful! API is working.'
            elif resp.status_code == 401:
                diagnostics['suggestion'] = 'Token is invalid or expired. Get a new token from your Tabbycat Change Password page.'
            elif resp.status_code == 403:
                diagnostics['suggestion'] = 'Access forbidden. Try Session Auth Fallback (admin username + password).'
            elif resp.status_code == 404:
                diagnostics['suggestion'] = f'Tournament slug "{self.slug}" not found.'
            else:
                diagnostics['suggestion'] = f'Unexpected status {resp.status_code}.'
        except Exception as e:
            diagnostics['steps'].append({
                'step': 'Tournament teams list',
                'status': 0,
                'ok': False,
                'error': str(e)
            })
            diagnostics['suggestion'] = f'Connection error: {str(e)}.'

        return diagnostics


    # ------------------------------------------------------------------
    # v4.0: read what already exists on the tab site (GET only)
    # ------------------------------------------------------------------
    def _get_list(self, url):
        """GET a list endpoint (follows pagination). Returns (items, error_message)."""
        items = []
        next_url = url
        for _ in range(200):
            try:
                time.sleep(0.2)
                resp = self.session.get(next_url, timeout=60)
            except requests.exceptions.RequestException as e:
                return None, str(e)[:150]
            if resp.status_code != 200:
                return None, f"HTTP {resp.status_code} on GET {next_url.replace(self.base_url, '')}"
            try:
                data = resp.json()
            except ValueError:
                return None, 'response was not JSON'
            if isinstance(data, list):
                items.extend(data)
                return items, None
            if isinstance(data, dict) and isinstance(data.get('results'), list):
                items.extend(data['results'])
                next_url = data.get('next')
                if not next_url:
                    return items, None
                continue
            return None, 'unexpected response shape'
        return items, None

    def load_existing(self, need_teams, need_adjudicators, need_speakers):
        """
        Read existing institutions / teams / adjudicators / speakers from the tab site.
        Returns (existing_dict, problems). If a needed list cannot be read we report a
        problem so the caller can stop instead of risking duplicates.
        """
        problems = []
        existing = {'institutions': [], 'teams': [], 'adjudicators': [], 'speakers': []}

        items, err = self._get_list(self._global_url('/institutions'))
        if items is None:
            items, err2 = self._get_list(self._tournament_url('/institutions'))
            if items is None:
                problems.append(f"institutions could not be read ({err})")
        existing['institutions'] = items or []

        if need_teams:
            items, err = self._get_list(self._tournament_url('/teams'))
            if items is None:
                problems.append(f"teams could not be read ({err})")
            existing['teams'] = items or []

        if need_adjudicators:
            items, err = self._get_list(self._tournament_url('/adjudicators'))
            if items is None:
                problems.append(f"adjudicators could not be read ({err})")
            existing['adjudicators'] = items or []

        if need_speakers:
            items, err = self._get_list(self._tournament_url('/speakers'))
            if items is not None:
                existing['speakers'] = items
            else:
                # Fall back to the speakers nested inside each team, if they carry names
                teams = existing['teams']
                nested_ok = bool(teams) and all(
                    isinstance(t.get('speakers'), list) and
                    all(isinstance(sp, dict) and sp.get('name') for sp in t['speakers'])
                    for t in teams)
                if nested_ok:
                    existing['speakers'] = []
                else:
                    problems.append(f"speakers could not be read ({err})")
        return existing, problems

    def create_institution(self, name, code, region=''):
        if code in self.created_institutions:
            return self.created_institutions[code]

        data = {"name": name, "code": code}
        if region:
            data["region"] = region
        url = self._global_url('/institutions')
        result = self._request('POST', url, data)

        if result and 'url' in result:
            self.created_institutions[code] = result['url']
            return result['url']
        elif result and 'id' in result:
            inst_url = f"{self.base_url}/api/v1/institutions/{result['id']}/"
            self.created_institutions[code] = inst_url
            return inst_url
        return None

    def create_team(self, institution_url, reference, short_reference,
                    use_institution_prefix=True, emoji='', speakers=None, code_name=''):
        data = {
            "reference": reference,
            "short_reference": short_reference or reference,
            "use_institution_prefix": use_institution_prefix,
        }
        if institution_url is not None:
            data["institution"] = institution_url
        if emoji:
            data["emoji"] = emoji
        if code_name:
            data["code_name"] = code_name
        if speakers:
            data["speakers"] = speakers

        url = self._tournament_url('/teams')
        return self._request('POST', url, data)

    def create_adjudicator(self, name, institution_url=None, email='',
                           gender='', base_score=None, independent=False,
                           adj_core=False, notes=''):
        data = {
            "name": name,
            "independent": independent,
            "adj_core": adj_core,
            "institution_conflicts": [],
            "team_conflicts": [],
            "adjudicator_conflicts": [],
        }

        if institution_url:
            data["institution"] = institution_url
        else:
            data["institution"] = None

        if email:
            data["email"] = email
        if gender:
            data["gender"] = gender
        if base_score is not None:
            data["base_score"] = base_score
        if notes:
            data["notes"] = notes

        url = self._tournament_url('/adjudicators')
        return self._request('POST', url, data)

    def get_speaker_categories(self):
        """Fetch existing speaker categories for the tournament."""
        url = self._tournament_url('/speaker-categories')
        return self._request('GET', url)

    def create_speaker_category(self, name):
        """Create a new speaker category."""
        data = {"name": name}
        url = self._tournament_url('/speaker-categories')
        return self._request('POST', url, data)

    def create_speaker(self, team_url, name, email='', gender='', categories=None):
        """
        v3.6 FIXES:
        - team_url: accepts the full team URL directly (stripped of trailing slashes
          to avoid DRF hyperlink resolver issues).
        - categories: accepts a list of category URLs to assign to the speaker.
        """
        # Strip trailing slashes to prevent DRF "Invalid hyperlink" errors
        team_url = team_url.rstrip('/') if team_url else team_url

        data = {
            "name": name,
            "team": team_url,
            "categories": list(categories) if categories else [],
        }
        if email:
            data["email"] = email
        if gender:
            data["gender"] = gender

        url = self._tournament_url('/speakers')
        return self._request('POST', url, data)


# =============================================================================
# CSV PROCESSORS
# =============================================================================

def process_institutions(rows):
    results = []
    errors = []
    seen_codes = set()
    for idx, row in enumerate(rows, start=2):
        name = clean_string(row.get('name', ''))
        code = clean_string(row.get('code', ''))
        if not name and not code:
            continue
        if not name:
            errors.append(f"Row {idx}: Missing institution name")
            continue
        if not code:
            errors.append(f"Row {idx}: Missing code for '{name}'")
            continue
        if code in seen_codes:
            errors.append(f"Row {idx}: Duplicate institution code '{code}'")
            continue
        seen_codes.add(code)
        results.append({
            'name': name,
            'code': code,
            'region': clean_string(row.get('region', ''))
        })
    return results, errors


def process_adjudicators(rows):
    results = []
    errors = []
    for idx, row in enumerate(rows, start=2):
        name = clean_string(row.get('name', ''))
        if not name:
            continue
        gender = clean_string(row.get('gender', ''))
        gender_norm = ''
        if gender.upper() in ['M', 'MALE']:
            gender_norm = 'M'
        elif gender.upper() in ['F', 'FEMALE']:
            gender_norm = 'F'
        elif gender.upper() in ['O', 'OTHER']:
            gender_norm = 'O'
        results.append({
            'name': name,
            'institution': clean_string(row.get('institution', '')),
            'email': clean_string(row.get('email', '')),
            'gender': gender_norm,
            'base_score': parse_float_or_none(row.get('base_score')),
            'independent': parse_bool(row.get('independent')),
            'adj_core': parse_bool(row.get('adj_core')),
            'notes': clean_string(row.get('notes', ''))
        })
    return results, errors


def process_teams(rows):
    results = []
    errors = []
    seen_refs = set()
    for idx, row in enumerate(rows, start=2):
        institution = clean_string(row.get('institution', ''))
        ref = clean_string(row.get('reference', ''))
        if not ref:
            errors.append(f"Row {idx}: Missing reference for institution '{institution or '(none)'}'")
            continue
        # Independent teams: allow blank institution
        key = f"{institution or '__INDEPENDENT__'}:{ref}"
        if key in seen_refs:
            errors.append(f"Row {idx}: Duplicate team reference '{ref}' for institution '{institution or '(none)'}'")
            continue
        seen_refs.add(key)
        results.append({
            'institution': institution,
            'reference': ref,
            'short_reference': clean_string(row.get('short_reference', ref)),
            'code_name': clean_string(row.get('code_name', '')),
            'use_institution_prefix': parse_bool(row.get('use_institution_prefix', True)),
            'emoji': clean_string(row.get('emoji', '')),
            'team_name_human': clean_string(row.get('team_name (human)', ''))
        })
    return results, errors


def process_speakers(rows, max_speakers=None, debate_format=None):
    results = []
    errors = []
    for idx, row in enumerate(rows, start=2):
        name = clean_string(row.get('name', ''))
        if not name:
            continue
        gender = clean_string(row.get('gender', ''))
        gender_norm = ''
        if gender.upper() in ['M', 'MALE']:
            gender_norm = 'M'
        elif gender.upper() in ['F', 'FEMALE']:
            gender_norm = 'F'
        elif gender.upper() in ['O', 'OTHER']:
            gender_norm = 'O'
        results.append({
            'name': name,
            'gender': gender_norm,
            'email': clean_string(row.get('email', '')),
            'phone': clean_string(row.get('phone', '')),
            'anonymous': parse_bool(row.get('anonymous')),
            'team': clean_string(row.get('team', '')),
            'categories': clean_string(row.get('categories', '')),
            'initials_match': clean_string(row.get('initials_match', ''))
        })

    if max_speakers and max_speakers > 0:
        label = FORMAT_LABELS.get(debate_format, 'selected')
        team_counts = {}
        filtered = []
        for spk in results:
            team = spk['team']
            team_counts[team] = team_counts.get(team, 0) + 1
            if team_counts[team] <= max_speakers:
                filtered.append(spk)
            elif team_counts[team] == max_speakers + 1:
                errors.append(f"Team '{team}': Only first {max_speakers} speakers imported ({label} format). Skipped extra speakers.")
        results = filtered

        # Warn (does not block) when a team has fewer speakers than the format expects
        min_speakers = FORMAT_MIN_SPEAKERS.get(debate_format)
        if min_speakers:
            for team, n in team_counts.items():
                if team and n < min_speakers:
                    errors.append(f"Team '{team}': only {n} speaker(s) found; {label} expects at least {min_speakers}.")

    return results, errors


# =============================================================================
# CSV GENERATORS
# =============================================================================

def generate_institutions_csv(institutions):
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(['name', 'code', 'region'])
    for inst in institutions:
        writer.writerow([inst['name'], inst['code'], inst.get('region', '')])
    return output.getvalue()


def generate_adjudicators_csv(adjudicators):
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(['institution', 'name', 'email', 'gender', 'base_score', 'independent', 'adj_core', 'notes'])
    for adj in adjudicators:
        writer.writerow([
            adj['institution'], adj['name'], adj['email'], adj['gender'],
            adj['base_score'] if adj['base_score'] is not None else '',
            'TRUE' if adj['independent'] else 'FALSE',
            'TRUE' if adj['adj_core'] else 'FALSE',
            adj['notes']
        ])
    return output.getvalue()


def generate_teams_csv(teams):
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(['institution', 'reference', 'short_reference', 'use_institution_prefix', 'emoji'])
    for team in teams:
        writer.writerow([
            team['institution'], team['reference'], team['short_reference'],
            'TRUE' if team['use_institution_prefix'] else 'FALSE',
            team['emoji']
        ])
    return output.getvalue()


def generate_speakers_csv(speakers):
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(['team', 'name', 'email', 'phone', 'gender', 'anonymous', 'categories'])
    for spk in speakers:
        writer.writerow([
            spk['team'], spk['name'], spk['email'], spk['phone'],
            spk['gender'], 'TRUE' if spk['anonymous'] else 'FALSE',
            spk.get('categories', '')
        ])
    return output.getvalue()


# =============================================================================
# BATCH-AWARE IMPORT (v4.0)
# =============================================================================

class ImportAbort(Exception):
    pass


def new_report():
    report = {key: {'created': [], 'existing': [], 'failed': []}
              for key in ('institutions', 'teams', 'adjudicators', 'speakers')}
    report['notes'] = []
    return report


def build_institution_index(items, base_url):
    """Index institutions already on the site by code and by name."""
    index = {'by_code': {}, 'by_name': {}, 'by_id': {}, 'by_key': {}}
    for item in items:
        if not isinstance(item, dict):
            continue
        inst_id = item.get('id')
        url = item.get('url') or ''
        if inst_id is None:
            inst_id = url_id(url, 'institutions')
        if not url and inst_id is not None:
            url = f"{base_url}/api/v1/institutions/{inst_id}"
        entry = (inst_id, url)
        if norm(item.get('code')):
            index['by_code'][norm(item.get('code'))] = entry
        if norm(item.get('name')):
            index['by_name'][norm(item.get('name'))] = entry
        for text in (item.get('name'), item.get('code')):
            for key in (inst_key(text), inst_key(text, True)):
                if key:
                    index['by_key'].setdefault(key, entry)
        if inst_id is not None:
            index['by_id'][inst_id] = {'code': item.get('code') or '', 'name': item.get('name') or ''}
    return index


def find_existing_institution(index, code, name=''):
    """
    Match a CSV institution (code or name) to one on the site: first ignoring case / extra spaces,
    then ignoring punctuation too ("St. Stephen's College" == "St Stephens College").
    """
    for key, table in ((norm(code), index['by_code']), (norm(name), index['by_name']),
                       (norm(code), index['by_name']), (norm(name), index['by_code'])):
        if key and key in table:
            return table[key]
    for text in (name, code):
        for key in (inst_key(text), inst_key(text, True)):
            if key and key in index['by_key']:
                return index['by_key'][key]
    return None


def run_batch_import(api, institutions, teams, adjudicators, speakers, max_speakers, dry_run):
    """
    Compare the uploaded rows with what is already on the tab site and import only
    what is new. Returns (report, new_rows).
      report   -> per type: created / existing / failed lists (for the results page)
      new_rows -> the uploaded rows that were new (used for the "new entries only" CSVs)
    """
    report = new_report()
    new_rows = {'institutions': [], 'teams': [], 'adjudicators': [], 'speakers': []}

    existing, problems = api.load_existing(
        need_teams=bool(teams or speakers),
        need_adjudicators=bool(adjudicators),
        need_speakers=bool(speakers))
    if problems:
        raise ImportAbort('Could not read the existing entries from the tab site (' + '; '.join(problems) +
                          '). Nothing was imported, to avoid creating duplicates. '
                          'Check the API token has admin access and try again.')

    def failure_reason(errors_before):
        if len(api.stats['errors']) > errors_before:
            return api.stats['errors'][-1]
        return 'request failed'

    # ------------------------------------------------------------------ institutions
    inst_index = build_institution_index(existing['institutions'], api.base_url)
    resolved = {}   # norm(code) -> (id, url) for institutions in the uploaded file / already on site

    for inst in institutions:
        label = f"{inst['name']} ({inst['code']})"
        original_code = inst['code']
        if inst.get('auto_code'):
            # code was generated by the importer: match by NAME only, never grab a site institution by a made-up code
            found = find_existing_institution(inst_index, '', inst['name'])
        else:
            found = find_existing_institution(inst_index, inst['code'], inst['name'])
        if found:
            resolved[norm(original_code)] = found
            report['institutions']['existing'].append(label)
            continue
        if inst.get('auto_code'):
            code, n = inst['code'], 2
            while norm(code) in inst_index['by_code'] or norm(code) in resolved:
                code, n = f"{inst['code']}{n}", n + 1
            inst = dict(inst, code=code)
            label = f"{inst['name']} ({code})"
        if dry_run:
            resolved[norm(original_code)] = (f"new:{norm(inst['code'])}", f"new:{norm(inst['code'])}")
            resolved[norm(inst['code'])] = resolved[norm(original_code)]
            report['institutions']['created'].append(label)
            new_rows['institutions'].append(inst)
            continue
        before = len(api.stats['errors'])
        url = api.create_institution(inst['name'], inst['code'], inst.get('region', ''))
        if url:
            resolved[norm(original_code)] = (url_id(url, 'institutions'), url)
            resolved[norm(inst['code'])] = resolved[norm(original_code)]
            report['institutions']['created'].append(label)
            new_rows['institutions'].append(inst)
        else:
            report['institutions']['failed'].append((label, failure_reason(before)))

    def resolve_institution(code):
        """-> (inst_id, inst_url, status) where status is 'blank', 'ok' or 'missing'."""
        if not code:
            return None, None, 'blank'
        key = norm(code)
        if key in resolved:
            return resolved[key][0], resolved[key][1], 'ok'
        found = find_existing_institution(inst_index, code, code)
        if found:
            resolved[key] = found
            return found[0], found[1], 'ok'
        return None, None, 'missing'

    # ------------------------------------------------------------------ teams
    team_index = {}            # (institution id or None, norm(reference)) -> team url
    team_alias_site = {}       # norm(team name) -> team url  (teams already on the site)
    team_alias_csv = {}        # norm(team name) -> team url  (teams in the uploaded file)
    team_url_by_id = {}
    speaker_names = {}         # team url -> set of norm(speaker name) already attached

    for t in existing['teams']:
        if not isinstance(t, dict):
            continue
        t_id = t.get('id') if t.get('id') is not None else url_id(t.get('url'), 'teams')
        t_url = (t.get('url') or '').rstrip('/') or f"{api.base_url}/api/v1/tournaments/{api.slug}/teams/{t_id}"
        if t_id is not None:
            team_url_by_id[t_id] = t_url
        inst_id = url_id(t.get('institution'), 'institutions')
        team_index[(inst_id, norm(t.get('reference')))] = t_url
        inst_info = inst_index['by_id'].get(inst_id, {})
        names = {t.get('short_name'), t.get('long_name'), t.get('reference')}
        for prefix in (inst_info.get('code'), inst_info.get('name')):
            if prefix:
                names.add(f"{prefix} {t.get('reference')}")
                if t.get('short_reference'):
                    names.add(f"{prefix} {t.get('short_reference')}")
        for name in names:
            if norm(name):
                team_alias_site.setdefault(norm(name), t_url)
        speaker_names.setdefault(t_url, set())
        for sp in (t.get('speakers') or []):
            if isinstance(sp, dict) and sp.get('name'):
                speaker_names[t_url].add(norm(sp['name']))

    for sp in existing['speakers']:
        if not isinstance(sp, dict):
            continue
        t_ref = sp.get('team')
        t_id = url_id(t_ref, 'teams')
        t_url = team_url_by_id.get(t_id) or (str(t_ref).rstrip('/') if t_ref else None)
        if t_url and sp.get('name'):
            speaker_names.setdefault(t_url, set()).add(norm(sp['name']))

    for team in teams:
        code = team['institution']
        inst_id, inst_url, status = resolve_institution(code)
        default_name = f"{code or 'Independent'} {team['reference']}"
        label = default_name if not team['team_name_human'] else f"{team['team_name_human']} [{default_name}]"
        if status == 'missing' and team.get('inst_soft'):
            report['notes'].append(f"Team '{label}': institution '{code}' was not found on the tab site, "
                                   f"so the team is imported as independent.")
            inst_id, inst_url, status = None, None, 'blank'
        if status == 'missing':
            why = f"institution '{code}' was not found in your institutions file or on the tab site"
            if 'inst_soft' in team:      # team came from the combined Teams + Speakers file
                why += (" - add it to the Institutions file (name + code), or use the option "
                        "'Import the team as independent'")
            report['teams']['failed'].append((label, why))
            continue
        key = (inst_id if status == 'ok' else None, norm(team['reference']))
        site_url = team_index.get(key)
        names_for_team = {norm(default_name)}
        if team['team_name_human']:
            names_for_team.add(norm(team['team_name_human']))

        if site_url:
            for n in names_for_team:
                team_alias_csv[n] = site_url
            report['teams']['existing'].append(label)
            continue

        if dry_run:
            fake_url = f"new-team:{norm(default_name)}"
            for n in names_for_team:
                team_alias_csv[n] = fake_url
            speaker_names.setdefault(fake_url, set())
            report['teams']['created'].append(label)
            new_rows['teams'].append(team)
            continue

        before = len(api.stats['errors'])
        result = api.create_team(
            institution_url=inst_url,
            reference=team['reference'],
            short_reference=team['short_reference'],
            use_institution_prefix=team['use_institution_prefix'],
            emoji=team['emoji'],
            code_name=team['code_name'],
            speakers=None)
        if result and 'id' in result:
            t_url = (result.get('url') or '').rstrip('/') or f"{api.base_url}/api/v1/tournaments/{api.slug}/teams/{result['id']}"
            team_index[key] = t_url
            speaker_names.setdefault(t_url, set())
            for n in names_for_team:
                team_alias_csv[n] = t_url
            report['teams']['created'].append(label)
            new_rows['teams'].append(team)
        else:
            report['teams']['failed'].append((label, failure_reason(before)))

    # ------------------------------------------------------------------ speaker categories
    category_map = {}
    if speakers and not dry_run:
        existing_cats = api.get_speaker_categories()
        if isinstance(existing_cats, list):
            for cat in existing_cats:
                if isinstance(cat, dict) and 'name' in cat:
                    cat_url = (cat.get('url') or '').rstrip('/')
                    if not cat_url and 'id' in cat:
                        cat_url = f"{api.base_url}/api/v1/tournaments/{api.slug}/speaker-categories/{cat['id']}"
                    category_map[cat['name']] = cat_url

            needed = set()
            for spk in speakers:
                for cat_name in [c.strip() for c in (spk.get('categories') or '').split(',')]:
                    if cat_name and cat_name not in category_map:
                        needed.add(cat_name)
            for cat_name in needed:
                result = api.create_speaker_category(cat_name)
                if result:
                    cat_url = (result.get('url') or '').rstrip('/')
                    if not cat_url and 'id' in result:
                        cat_url = f"{api.base_url}/api/v1/tournaments/{api.slug}/speaker-categories/{result['id']}"
                    category_map[cat_name] = cat_url

    # ------------------------------------------------------------------ speakers
    for spk in speakers:
        label = f"{spk['name']} ({spk['team']})"
        t_key = norm(spk['team'])
        t_url = team_alias_csv.get(t_key) or team_alias_site.get(t_key)
        if not t_url:
            report['speakers']['failed'].append((label, f"team '{spk['team']}' was not found in your teams file or on the tab site"))
            continue
        names = speaker_names.setdefault(t_url, set())
        if norm(spk['name']) in names:
            report['speakers']['existing'].append(label)
            continue
        if max_speakers and len(names) >= max_speakers:
            report['speakers']['failed'].append((label, f"team already has {len(names)} speakers (limit for this format is {max_speakers})"))
            continue

        if dry_run:
            names.add(norm(spk['name']))
            report['speakers']['created'].append(label)
            new_rows['speakers'].append(spk)
            continue

        cat_urls = []
        for cat_name in [c.strip() for c in (spk.get('categories') or '').split(',')]:
            if cat_name in category_map:
                cat_urls.append(category_map[cat_name])
        before = len(api.stats['errors'])
        result = api.create_speaker(team_url=t_url, name=spk['name'], email=spk.get('email', ''),
                                    gender=spk.get('gender', ''), categories=cat_urls)
        if result:
            names.add(norm(spk['name']))
            report['speakers']['created'].append(label)
            new_rows['speakers'].append(spk)
        else:
            report['speakers']['failed'].append((label, failure_reason(before)))

    # ------------------------------------------------------------------ adjudicators
    adj_keys = set()
    adj_emails = set()
    for a in existing['adjudicators']:
        if not isinstance(a, dict):
            continue
        adj_keys.add((norm(a.get('name')), url_id(a.get('institution'), 'institutions')))
        if norm(a.get('email')):
            adj_emails.add(norm(a.get('email')))

    for adj in adjudicators:
        code = adj['institution']
        inst_id, inst_url, status = resolve_institution(code)
        label = f"{adj['name']} ({code})" if code else f"{adj['name']} (independent)"
        if status == 'missing':
            report['adjudicators']['failed'].append(
                (label, f"institution '{code}' was not found in your institutions file or on the tab site"))
            continue
        key = (norm(adj['name']), inst_id if status == 'ok' else None)
        email_key = norm(adj.get('email'))
        if key in adj_keys or (email_key and email_key in adj_emails):
            report['adjudicators']['existing'].append(label)
            continue

        if dry_run:
            adj_keys.add(key)
            if email_key:
                adj_emails.add(email_key)
            report['adjudicators']['created'].append(label)
            new_rows['adjudicators'].append(adj)
            continue

        before = len(api.stats['errors'])
        result = api.create_adjudicator(
            name=adj['name'], institution_url=inst_url, email=adj['email'], gender=adj['gender'],
            base_score=adj['base_score'], independent=adj['independent'],
            adj_core=adj['adj_core'], notes=adj['notes'])
        if result:
            adj_keys.add(key)
            if email_key:
                adj_emails.add(email_key)
            report['adjudicators']['created'].append(label)
            new_rows['adjudicators'].append(adj)
        else:
            report['adjudicators']['failed'].append((label, failure_reason(before)))

    return report, new_rows


# =============================================================================
# FLASK ROUTES
# =============================================================================

@app.route('/')
def index():
    return render_template('index.html')


@app.route('/test-connection', methods=['POST'])
def test_connection():
    data = request.get_json()
    api = TabbycatAPI(
        data.get('base_url', ''),
        data.get('token', ''),
        data.get('slug', ''),
        username=data.get('username'),
        password=data.get('password')
    )
    diagnostics = api.test_connection()
    return jsonify(diagnostics)


@app.route('/existing', methods=['POST'])
def existing_summary():
    """Show how many institutions / teams / judges / speakers are already on the tab site."""
    data = request.get_json() or {}
    try:
        api = TabbycatAPI(data.get('base_url', ''), data.get('token', ''), data.get('slug', ''),
                          username=data.get('username') or None, password=data.get('password') or None)
        existing, problems = api.load_existing(True, True, True)
        counts = {
            'institutions': len(existing['institutions']),
            'teams': len(existing['teams']),
            'adjudicators': len(existing['adjudicators']),
            'speakers': len(existing['speakers']) or sum(len(t.get('speakers') or []) for t in existing['teams']
                                                         if isinstance(t, dict)),
        }
        return jsonify({'ok': not problems, 'counts': counts, 'problems': problems})
    except Exception as e:
        return jsonify({'ok': False, 'counts': {}, 'problems': [str(e)]})


@app.route('/api-diagnose', methods=['POST'])
def api_diagnose():
    data = request.get_json()
    base_url = data.get('base_url', '').rstrip('/')
    token = data.get('token', '').strip()
    slug = data.get('slug', '').strip('/')

    results = []
    session = requests.Session()
    session.headers.update({
        'User-Agent': 'TabbycatImporter/4.0 (Diagnostic)',
        'Accept': 'application/json'
    })
    if token:
        session.headers['Authorization'] = f'Token {token}'

    paths_to_try = [
        ('GET', '/api/v1/institutions'),
        ('GET', f'/api/v1/tournaments/{slug}/institutions'),
        ('GET', f'/api/v1/tournaments/{slug}/teams'),
        ('GET', f'/api/v1/tournaments/{slug}/adjudicators'),
        ('GET', '/api/v1/'),
        ('GET', '/api/'),
    ]

    for method, path in paths_to_try:
        url = f"{base_url}{path}"
        try:
            resp = session.get(url, timeout=10)
            results.append({
                'method': method,
                'url': url,
                'status': resp.status_code,
                'body_preview': resp.text[:150] if resp.text else '(empty)'
            })
        except Exception as e:
            results.append({
                'method': method,
                'url': url,
                'status': 0,
                'error': str(e)
            })

    return jsonify({'results': results})


def read_optional_upload(field):
    f = request.files.get(field)
    if f and f.filename:
        return read_uploaded_file(f)
    return None


@app.route('/upload', methods=['POST'])
def upload():
    mode = request.form.get('mode', 'csv')
    debate_format = request.form.get('debate_format', 'bp')
    max_speakers = FORMAT_MAX_SPEAKERS.get(debate_format, 3)
    dry_run = request.form.get('dry_run') == 'on'

    try:
        inst_rows = read_optional_upload('institutions')
        adj_rows = read_optional_upload('adjudicators')
        team_rows = read_optional_upload('teams')
        speaker_rows = read_optional_upload('speakers')
        combined_rows = read_optional_upload('teams_speakers')

        if all(rows is None for rows in (inst_rows, adj_rows, team_rows, speaker_rows, combined_rows)):
            flash('Please upload at least one file (institutions, adjudicators, teams, speakers or the combined teams + speakers file).', 'error')
            return redirect(url_for('index'))

        institutions, inst_errors = process_institutions(inst_rows) if inst_rows is not None else ([], [])
        adjudicators, adj_errors = process_adjudicators(adj_rows) if adj_rows is not None else ([], [])
        teams, team_errors = process_teams(team_rows) if team_rows is not None else ([], [])
        speakers, speaker_errors = (process_speakers(speaker_rows, max_speakers=max_speakers, debate_format=debate_format)
                                    if speaker_rows is not None else ([], []))

        # ---- combined "Teams + Speakers" file: split every row into a team + its speakers (+ institutions)
        combined = None
        if combined_rows is not None:
            try:
                combined = parse_combined(
                    combined_rows,
                    options={
                        'institution_policy': request.form.get('institution_policy', 'create'),
                        'prefix': request.form.get('team_prefix', 'auto'),
                        'novice_category': request.form.get('novice_category') == 'on',
                    },
                    known_institutions=institutions,
                    max_speakers=max_speakers, debate_format=debate_format,
                    min_speakers=FORMAT_MIN_SPEAKERS.get(debate_format))
            except ValueError as e:
                flash(f'Combined teams + speakers file: {e}', 'error')
                return redirect(url_for('index'))
            known_codes = {norm(i['code']) for i in institutions}
            institutions = institutions + [i for i in combined['institutions'] if norm(i['code']) not in known_codes]
            teams = teams + combined['teams']
            speakers = speakers + combined['speakers']

        # CSV-only mode has no tab site to look at: check institution codes against the file, if there is one.
        if mode != 'api' and institutions:
            institution_codes = {i['code'] for i in institutions}
            for t in teams:
                if t.get('inst_soft'):
                    continue
                if t['institution'] and t['institution'] not in institution_codes:
                    team_errors.append(f"Team '{t['institution']} {t['reference']}': institution code '{t['institution']}' not found in institutions file")

        api_results = None
        report = None
        new_rows = None

        if mode == 'api':
            base_url = request.form.get('api_url', '').strip()
            token = request.form.get('api_token', '').strip()
            slug = request.form.get('tournament_slug', '').strip()
            username = request.form.get('api_username', '').strip() or None
            password = request.form.get('api_password', '').strip() or None

            if not all([base_url, token, slug]):
                flash('API URL, Token, and Tournament Slug are required for API mode', 'error')
                return redirect(url_for('index'))

            api = TabbycatAPI(base_url, token, slug, username=username, password=password)
            diagnostics = api.test_connection()

            if not diagnostics['ok']:
                flash(f"API Connection Failed: {diagnostics['suggestion']}", 'error')
                for step in diagnostics['steps']:
                    flash(f"  {step['step']}: HTTP {step.get('status', 'ERR')}", 'info')
                return redirect(url_for('index'))

            try:
                report, new_rows = run_batch_import(api, institutions, teams, adjudicators, speakers,
                                                    max_speakers, dry_run)
            except ImportAbort as e:
                flash(str(e), 'error')
                return redirect(url_for('index'))

            api_results = api.stats

        # CSV downloads: in API mode only the NEW entries (handy for a manual import of this batch)
        if new_rows is not None:
            csv_institutions, csv_teams = new_rows['institutions'], new_rows['teams']
            csv_adjudicators, csv_speakers = new_rows['adjudicators'], new_rows['speakers']
        else:
            csv_institutions, csv_teams, csv_adjudicators, csv_speakers = institutions, teams, adjudicators, speakers

        download_id = uuid.uuid4().hex
        DOWNLOAD_CACHE[download_id] = {
            'teams_speakers': combined['clean_csv'] if combined else '',
            'institutions': generate_institutions_csv(csv_institutions),
            'adjudicators': generate_adjudicators_csv(csv_adjudicators),
            'teams': generate_teams_csv(csv_teams),
            'speakers': generate_speakers_csv(csv_speakers),
        }
        while len(DOWNLOAD_CACHE) > MAX_CACHED_DOWNLOADS:
            DOWNLOAD_CACHE.popitem(last=False)

        counts = {
            'inst': len(csv_institutions), 'adj': len(csv_adjudicators),
            'team': len(csv_teams), 'speaker': len(csv_speakers),
        }

        return render_template('results.html',
                               institutions=institutions,
                               inst_errors=inst_errors,
                               adjudicators=adjudicators,
                               adj_errors=adj_errors,
                               teams=teams,
                               team_errors=team_errors,
                               speakers=speakers,
                               speaker_errors=speaker_errors,
                               inst_count=len(institutions),
                               adj_count=len(adjudicators),
                               team_count=len(teams),
                               speaker_count=len(speakers),
                               csv_counts=counts,
                               combined=combined,
                               download_id=download_id,
                               mode=mode,
                               dry_run=dry_run,
                               report=report,
                               debate_format=debate_format,
                               api_results=api_results)

    except Exception as e:
        flash(f'Error: {str(e)}', 'error')
        return redirect(url_for('index'))


@app.route('/template/teams-speakers')
def teams_speakers_template():
    buffer = io.BytesIO(template_csv().encode('utf-8'))
    return send_file(buffer, mimetype='text/csv', as_attachment=True, download_name='teams_speakers_template.csv')


@app.route('/download/<download_id>/<file_type>')
def download(download_id, file_type):
    bundle = DOWNLOAD_CACHE.get(download_id)
    if not bundle or file_type not in bundle:
        flash('That download has expired. Please run the import again.', 'error')
        return redirect(url_for('index'))
    buffer = io.BytesIO(bundle[file_type].encode('utf-8'))
    return send_file(buffer, mimetype='text/csv', as_attachment=True, download_name=f'{file_type}.csv')


if __name__ == '__main__':
    port = int(os.environ.get('PORT', 5000))
    app.run(host='0.0.0.0', port=port, debug=False)
